"""Flag departures from the academic register required by CLAUDE.md.

The rule is stated in the project instructions but was applied inconsistently,
twice, which is the argument for making it checkable rather than remembered.
The patterns below are drawn from actual departures found in this repository
rather than from a general stylebook, so a hit is evidence of the failure mode
that has already occurred here.

The check is advisory. Some hits are legitimate, a quotation or a proper name
among them, and the script reports rather than enforces.

Usage:
    python check_register.py
    python check_register.py docs/one_file.md
    python check_register.py --quiet          # exit status only
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

# Each entry is a label and a pattern. Ordered by how badly each reads in a
# document intended for supervisors and for a manuscript.
#
# A pattern for the antithetical title, of the form "Attribution, not
# verification", was written and then removed. It matched seven ordinary
# contrastive sentences for every rhetorical title, and a check that reports
# mostly noise is one that stops being read. That failure mode is caught by
# reading rather than by regular expression.
CHECKS: list[tuple[str, str]] = [
    ("first or second person",
     r"\b(we|our|ours|us|you|your|yours|I'm|I've|let's)\b"),
    ("em dash",
     r"—"),
    ("rhetorical aside",
     r"\b(uncomfortable|worked around|is thin|poor use|straightforward|"
     r"proper place|to their advantage|ground is open|open ground|read first|"
     r"no stakeholder|worth noting|it is worth|the trick|the catch|"
     r"the headline|the finding that matters|pays off|big deal|"
     r"where it belongs|better asked|already found|the obvious question|"
     r"what happens next|what we need)\b"),
    ("imperative or address to the reader",
     r"(^|\. )(Look to|Note that|Consider|Remember|Bear in mind|Keep in mind|"
     r"Ask |Check |Use |Read )"),
    ("meta-commentary on the document",
     r"\b(this section exists|as stated above|as mentioned|"
     r"as we (saw|noted)|in this document I)\b"),
    ("hyperbole",
     r"\b(dissolves|transformative|catastrophic|fatal|devastating|"
     r"enormous|huge|massive|dramatic(ally)?|stunning)\b"),
    ("judgemental section heading",
     r"^#{1,4} .*\b(why .* is (weak|wrong|bad)|what went wrong|the problem with|"
     r"what has already been done)\b"),
    # Anchored to the start of a line, since a rhetorical label is a heading.
    # The same words inside a sentence are ordinary prose.
    ("rhetorical question as a label",
     r"^\s*#*\s*\**(Why (not|this|each|it)\b|Why .{0,24}matters|"
     r"(How|What|Whether) .{0,60}\?\s*\**\s*$)"),
    # Added after a fourth lapse. The three below were the actual forms it took:
    # evaluative words that sell rather than state, sentences about the document
    # instead of about the subject, and labels built on an absence.
    # "compelling" is qualified by a preceding determiner or copula, since as a
    # verb ("no deadline compelling submission") it is ordinary prose.
    ("promotional or evaluative wording",
     r"\b(differentiator|(is|was|a|an|the|most|very) compelling|powerful|"
     r"invaluable|remarkable|striking|crucially|importantly|elegant|"
     r"worth (doing|computing|pursuing|putting|having)|sharpens|"
     r"is a strength|stands out|shows promise)\b"),
    # "which is itself" alone matched ordinary prose about the subject, so the
    # pattern names the evaluative nouns that follow it in the failure cases.
    ("commentary on the document rather than the subject",
     r"\b(this (slide|deck|table|figure) (is|shows|exists|covers)|"
     r"the purpose of this|recorded here (is|are)|"
     r"which is itself a (caution|reminder|warning|sign)|"
     r"as this document|the present (slide|deck))\b"),
    ("label built on a negation or a contrast",
     r"^\s*#*\s*\**(What it does not|What (this|it) is not|Not (a|an|the)\b|"
     r"The (problem|trouble|difficulty) with)\b"),
    # A single ongoing effort is described by what it found, not by which sweep
    # of it did the finding. This was a specific correction.
    ("narration of the work's own passes",
     r"\b(second pass|first pass|second sweep|second look|"
     r"re-?read(ing)? of the literature|revisiting the literature)\b"),
]

SKIP_DIRS = {".git", "data", "figures", "model", "__pycache__",
             ".ipynb_checkpoints", "node_modules"}

# Prose reaches the reader through more than Markdown. A deck is read by the
# host group and a generator script carries the deck's every sentence, so both
# are checked. Omitting them is how the rule was missed a third time.
SUFFIXES = {".md", ".py", ".js", ".pptx", ".docx"}

# An instruction file directs an agent and is legitimately imperative, so the
# imperative check is not applied to it. The register rules still are.
IMPERATIVE_EXEMPT = {"CLAUDE.md"}


def gather(targets: list[str]) -> list[Path]:
    if targets:
        return [Path(t) for t in targets]
    out: list[Path] = []
    for p in sorted(ROOT.rglob("*")):
        if any(part in SKIP_DIRS for part in p.parts):
            continue
        # Office writes a lock file beside an open document. It is not
        # readable and is not content.
        if p.name.startswith("~$"):
            continue
        if p.suffix in SUFFIXES and p.is_file() and p.name != Path(__file__).name:
            out.append(p)
    return out


def _pptx_lines(path: Path) -> list[str]:
    """Slide text and speaker notes, one shape or note per line.

    Speaker notes are included deliberately. They are read aloud and they drift
    into a conversational register more readily than the slides do.
    """
    try:
        from pptx import Presentation
    except ImportError:
        return ["python-pptx unavailable, deck not checked"]
    lines = []
    for n, slide in enumerate(Presentation(path).slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                for part in shape.text_frame.text.split("\n"):
                    lines.append(f"[slide {n}] {part}")
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                lines.append(f"[slide {n} notes] {note}")
    return lines


def _docx_lines(path: Path) -> list[str]:
    import xml.etree.ElementTree as ET
    import zipfile
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    try:
        root = ET.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    except (OSError, KeyError, ET.ParseError):
        return []
    return [t for t in ("".join(x.text or "" for x in p.iter(W + "t")).strip()
                        for p in root.iter(W + "p")) if t]


def scan(path: Path) -> list[tuple[int, str, str]]:
    hits: list[tuple[int, str, str]] = []
    if path.suffix == ".pptx":
        lines = _pptx_lines(path)
    elif path.suffix == ".docx":
        lines = _docx_lines(path)
    else:
        try:
            lines = path.read_text(encoding="utf-8").splitlines()
        except (UnicodeDecodeError, OSError):
            return hits
    text = "\n".join(lines)
    in_fence = False
    for n, line in enumerate(text.splitlines(), 1):
        if line.lstrip().startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        # A markdown link target or a bare URL is not prose.
        prose = re.sub(r"\]\([^)]*\)|https?://\S+", "", line)
        for label, pattern in CHECKS:
            # In a heading the same word is usually a noun, as in "Use in the
            # present work", so the imperative check does not apply there.
            if label.startswith("imperative") and (
                    prose.lstrip().startswith("#") or path.name in IMPERATIVE_EXEMPT):
                continue
            flags = re.IGNORECASE | (re.MULTILINE if pattern.startswith("^") else 0)
            m = re.search(pattern, prose, flags)
            if m:
                frag = prose.strip()
                if len(frag) > 110:
                    start = max(0, m.start() - 45)
                    frag = "…" + frag[start:m.end() + 45].strip() + "…"
                hits.append((n, label, frag))
    return hits


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("targets", nargs="*", help="files to check, default all")
    ap.add_argument("--quiet", action="store_true", help="exit status only")
    args = ap.parse_args(argv)

    total = 0
    for path in gather(args.targets):
        hits = scan(path)
        if not hits:
            continue
        total += len(hits)
        if args.quiet:
            continue
        try:
            shown = path.relative_to(ROOT)
        except ValueError:
            shown = path
        print(f"\n{shown}")
        for n, label, frag in hits:
            print(f"  {n:>4}  {label:<38} {frag}")

    if not args.quiet:
        print(f"\n{total} advisory hit{'s' if total != 1 else ''}. "
              f"Some are legitimate; the check reports rather than enforces.")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())
